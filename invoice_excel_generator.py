"""
Invoice Excel Generator — fills the SmartInvoice AI Excel template
with invoice data from the database and embeds the QR code.

Usage (imported):
    from invoice_excel_generator import generate_invoice_excel
    xlsx_bytes = generate_invoice_excel(invoice_id)
    # Returns bytes; write to file or send as HTTP response

Usage (standalone):
    python invoice_excel_generator.py <InvoiceId>
"""
import io
import os
import pyodbc
import urllib.request
import tempfile
from pathlib import Path
from copy import copy as _copy
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font
from PIL import Image as PILImage

# ---- Config -----------------------------------------------------------------
SQL_SERVER = "localhost"
SQL_DRIVER = "ODBC Driver 17 for SQL Server"
DATABASE = "CompanyInvoiceDB"
PORTAL_DB = "MyInvoisPortal"

TEMPLATE_PATH = r"C:\Users\User\Downloads\SmartInvoice_AI_Invoice_Template.xlsx"


def get_conn():
    return pyodbc.connect(
        f"DRIVER={{{SQL_DRIVER}}};SERVER={SQL_SERVER};DATABASE={DATABASE};"
        f"Trusted_Connection=yes;TrustServerCertificate=yes;",
        timeout=15,
    )


def fetch_invoice_full(invoice_id):
    """Fetch a single invoice with all fields from both tables."""
    with get_conn() as conn:
        c = conn.cursor()
        c.execute(f"""
            SELECT
                i.InvoiceId, i.InvoiceNumber,
                i.SupplierName, i.SupplierTIN,
                i.SupplierAddress, i.SupplierContact,
                i.BuyerName, i.BuyerTIN, i.Mykad,
                i.BuyerAddress, i.BuyerContact,
                i.InvoiceDate, i.InvoiceTime,
                i.Subtotal, i.SSTAmount, i.TotalPayable,
                ii.Status, ii.UUID, ii.QRUrl, ii.VerifyUrl
            FROM dbo.Invoices i
            INNER JOIN [{PORTAL_DB}].dbo.InvoiceInfo ii ON ii.InvoiceId = i.InvoiceId
            WHERE i.InvoiceId = ?
        """, invoice_id)
        cols = [d[0] for d in c.description]
        row = c.fetchone()
        if row:
            return dict(zip(cols, row))
        return None


def _fmt_date(d):
    """Format date as YYYY-MM-DD."""
    if d is None:
        return ""
    s = str(d)
    if len(s) >= 10:
        return s[:10]
    return s


def _fmt_time(t):
    """Format time as HH:MM:SS."""
    if t is None:
        return ""
    s = str(t)
    if len(s) >= 8:
        return s[:8]
    return s


def _fmt_money(val):
    """Format a numeric value as float for Excel."""
    if val is None:
        return 0
    try:
        return float(val)
    except (ValueError, TypeError):
        return 0


def _download_qr_image(qr_url):
    """Download QR code image from URL. Returns PIL Image or None."""
    if not qr_url:
        return None
    try:
        req = urllib.request.Request(
            qr_url,
            headers={"User-Agent": "Mozilla/5.0"},
        )
        with urllib.request.urlopen(req, timeout=15) as resp:
            img_data = resp.read()
        return PILImage.open(io.BytesIO(img_data))
    except Exception:
        return None


def _auto_row_height(ws, row_num, min_height=18, line_height=15):
    """
    Auto-fit row height based on the longest text in that row.
    Considers merged-cell widths and wrap_text.
    """
    max_lines = 1
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=row_num, column=col_idx)
        val = cell.value
        if val is None or not isinstance(val, str):
            continue

        # Determine available width (consider merged cells)
        col_letter = cell.column_letter
        merged_ranges = [
            mc for mc in ws.merged_cells.ranges
            if row_num >= mc.min_row and row_num <= mc.max_row
            and col_idx >= mc.min_col and col_idx <= mc.max_col
        ]

        if merged_ranges:
            # Sum widths of all columns in the merge
            total_width = 0
            mc = merged_ranges[0]
            for c_idx in range(mc.min_col, mc.max_col + 1):
                from openpyxl.utils import get_column_letter
                cl = get_column_letter(c_idx)
                dim = ws.column_dimensions.get(cl)
                total_width += dim.width if dim and dim.width else 8.43
        else:
            dim = ws.column_dimensions.get(col_letter)
            total_width = dim.width if dim and dim.width else 8.43

        # Estimate chars per line (rough: width * 1.0 for 11pt font)
        chars_per_line = max(1, int(total_width))
        text_len = len(str(val))

        # Check if wrap_text is enabled
        wrap = cell.alignment.wrap_text if cell.alignment else False
        if wrap or text_len > chars_per_line:
            lines = max(1, -(-text_len // chars_per_line))  # ceiling division
        else:
            lines = 1

        # Also count explicit newlines
        lines = max(lines, str(val).count("\n") + 1)

        if lines > max_lines:
            max_lines = lines

    calculated = max(min_height, max_lines * line_height + 3)
    ws.row_dimensions[row_num].height = calculated


def generate_invoice_excel(invoice_id):
    """
    Generate an Excel invoice from database data for the given InvoiceId.
    Uses the SmartInvoice AI template, fills in data, embeds QR code.
    Returns the XLSX as bytes.
    Raises ValueError if the invoice is not found.
    """
    inv = fetch_invoice_full(invoice_id)
    if inv is None:
        raise ValueError(f"Invoice with Id={invoice_id} not found in database.")

    # Load the template
    wb = load_workbook(TEMPLATE_PATH)
    ws = wb["Invoice Template"]

    # Helper alignment objects
    wrap_top = Alignment(wrap_text=True, vertical="top")
    wrap_center = Alignment(wrap_text=True, vertical="center")
    left_top = Alignment(wrap_text=True, vertical="top", horizontal="left")

    # ---- SELLER DETAILS (rows 5-8) ----
    # Template: A5="Company Name" B5:C5 merged, A6="Supplier Address" B6:C6 merged,
    #           A7="Supplier Email / Phone" B7:C7 merged, A8="Supplier TIN" B8:C8 merged
    ws["B5"] = inv.get("SupplierName") or ""
    ws["B5"].alignment = wrap_top

    ws["B6"] = inv.get("SupplierAddress") or ""
    ws["B6"].alignment = wrap_top

    ws["B7"] = inv.get("SupplierContact") or ""
    ws["B7"].alignment = wrap_top

    ws["B8"] = inv.get("SupplierTIN") or ""
    ws["B8"].alignment = wrap_top

    # ---- INVOICE DETAILS (E5-E7) ----
    # Template: D5="Invoice No." E5=data, D6="Invoice Date" E6=data, D7="Time" E7=data
    ws["E5"] = inv.get("InvoiceNumber") or ""
    ws["E5"].alignment = wrap_center

    ws["E6"] = _fmt_date(inv.get("InvoiceDate"))
    ws["E6"].alignment = wrap_center

    ws["E7"] = _fmt_time(inv.get("InvoiceTime"))
    ws["E7"].alignment = wrap_center

    # ---- BILL TO / BUYER DETAILS (rows 11-15) ----
    # Template: A11="Customer Name" B11:E11 merged, A12="Buyer Address" B12:E12 merged,
    #           A13="MyKad" B13:E13 merged, A14="Buyer Email / Phone" B14 (NOT merged),
    #           A15="Buyer TIN" B15:E15 merged
    ws["B11"] = inv.get("BuyerName") or ""
    ws["B11"].alignment = wrap_top

    ws["B12"] = inv.get("BuyerAddress") or ""
    ws["B12"].alignment = wrap_top

    ws["B13"] = inv.get("Mykad") or ""
    ws["B13"].alignment = wrap_top

    ws["B14"] = inv.get("BuyerContact") or ""
    ws["B14"].alignment = wrap_top

    ws["B15"] = inv.get("BuyerTIN") or ""
    ws["B15"].alignment = wrap_top

    # ---- Item table (row 18 = first item, clear rows 19-22) ----
    # Template: row 18 has A18=No, B18=Description, C18=Qty, D18=Unit Price, E18=Amount(=C18*D18)
    subtotal = _fmt_money(inv.get("Subtotal"))
    sst_amount = _fmt_money(inv.get("SSTAmount"))

    ws["A18"] = 1
    ws["B18"] = "Invoice total amount"
    ws["C18"] = 1
    ws["D18"] = subtotal
    # E18 already has formula =C18*D18 — keep it

    # Clear rows 19-22 (item rows 2-5)
    for row_num in range(19, 23):
        for col in "ABCDE":
            ws[f"{col}{row_num}"] = None

    # ---- Summary section (rows 24-28) ----
    # Template: E24==SUM(E18:E22) Subtotal, E25=Discount, E26=Tax, E27=Other charges,
    #           E28==E24-E25+E26+E27 Total
    # E24 has formula =SUM(E18:E22) — keep it (auto-calculates from item amounts)
    ws["E25"] = 0                                        # Discount
    ws["E26"] = sst_amount                               # Tax (SST)
    ws["E27"] = 0                                        # Other charges
    # E28 has formula =E24-E25+E26+E27 — keep it (auto-calculates total)

    # ---- Auto-fit row heights for data rows ----
    # Supplier section rows 5-8
    for r in [5, 6, 7, 8]:
        _auto_row_height(ws, r, min_height=20, line_height=15)

    # Buyer section rows 11-15
    for r in [11, 12, 13, 14, 15]:
        _auto_row_height(ws, r, min_height=20, line_height=15)

    # Item row 18
    _auto_row_height(ws, 18, min_height=20, line_height=15)

    # ---- QR Code image ----
    qr_url = inv.get("QRUrl") or ""
    status = inv.get("Status") or ""

    if status == "submitted" and qr_url:
        pil_img = _download_qr_image(qr_url)
        if pil_img is not None:
            # Resize to reasonable size for Excel
            pil_img = pil_img.resize((200, 200), PILImage.LANCZOS)
            # Save to temp file for openpyxl
            tmp_path = os.path.join(
                tempfile.gettempdir(),
                f"qr_{invoice_id}.png"
            )
            pil_img.save(tmp_path, "PNG")

            # Set column widths for QR area (columns F and G)
            ws.column_dimensions["F"].width = 4     # narrow spacer
            ws.column_dimensions["G"].width = 28    # QR code area

            # Embed QR image at G4 (top-right area)
            xl_img = XLImage(tmp_path)
            xl_img.width = 200
            xl_img.height = 200
            ws.add_image(xl_img, "G4")

            # Add "Scan to Verify" label at G2 (above the QR code)
            ws["G2"] = "Scan to Verify"
            ws["G2"].font = Font(bold=True, size=11)
            ws["G2"].alignment = Alignment(horizontal="center", vertical="center")

            # Add UUID label and value below QR code area
            ws["G17"] = "Validation UUID:"
            ws["G17"].font = Font(bold=True, size=9)
            ws["G17"].alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)

            ws["G18"] = inv.get("UUID") or ""
            ws["G18"].font = Font(size=9)
            ws["G18"].alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)

    # ---- Save to bytes ----
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ---- CLI --------------------------------------------------------------------
if __name__ == "__main__":
    import sys

    if len(sys.argv) != 2:
        print("Usage: python invoice_excel_generator.py <InvoiceId>")
        sys.exit(1)

    iid = int(sys.argv[1])
    print(f"Generating Excel for InvoiceId={iid}...")

    inv = fetch_invoice_full(iid)
    if inv is None:
        print(f"ERROR: Invoice Id={iid} not found.")
        sys.exit(1)

    xlsx_bytes = generate_invoice_excel(iid)
    filename = f"Invoice_{iid}_{inv.get('InvoiceNumber', 'unknown')}.xlsx"
    out_path = Path(__file__).parent / filename
    out_path.write_bytes(xlsx_bytes)

    print(f"Generated: {out_path}")
    print(f"  Invoice: {inv.get('InvoiceNumber')}")
    print(f"  Supplier: {inv.get('SupplierName')} (TIN: {inv.get('SupplierTIN')})")
    print(f"  Date: {inv.get('InvoiceDate')} Time: {inv.get('InvoiceTime')}")
    print(f"  Total: RM {float(inv.get('TotalPayable') or 0):,.2f}")
    print(f"  Status: {inv.get('Status')}")
    if inv.get("UUID"):
        print(f"  UUID: {inv['UUID']}")
    if inv.get("QRUrl"):
        print(f"  QR URL: {inv['QRUrl'][:80]}...")
    print(f"  File size: {len(xlsx_bytes)} bytes")
